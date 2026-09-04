import csv
import io
import json
import re
from datetime import date, timedelta
from difflib import SequenceMatcher
from io import BytesIO, StringIO

from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader


# LABELS RECONNUS DANS LES DOCUMENTS

LABELS = {
    "titre": ["titre", "title", "sujet", "intitulé", "intitule"],
    "description": ["description", "résumé", "resume", "détails", "details", "mission"],
    "entreprise": ["entreprise", "société", "societe", "company", "organisme"],
    "technologies": ["technologies", "compétences", "competences", "stack", "outils"],
    "duree": ["durée", "duree", "duration"],
    "localisation": ["localisation", "lieu", "ville", "location", "site"],
    "niveau_requis": ["niveau", "niveau requis", "niveau d'étude", "niveau d'etude", "level"],
    "domaine": ["domaine", "département", "departement", "catégorie", "categorie"],
}

CHAMPS_TEXTE = ["titre", "description", "entreprise", "technologies", "duree", "localisation", "niveau_requis"]

# Nombre de jours au-delà duquel un sujet toujours "Disponible" est
# considéré comme potentiellement obsolète.
SEUIL_OBSOLESCENCE_JOURS = 90


def _normaliser(texte: str) -> str:
    return re.sub(r"\s+", " ", (texte or "")).strip()


def _detecter_label(ligne: str):
    """Si la ligne suit le format 'Label : valeur', renvoie (champ, valeur)."""

    match = re.match(r"^\s*[-*•]?\s*([A-Za-zÀ-ÿ' ]{2,30})\s*[:\-]\s*(.+)$", ligne)

    if not match:
        return None

    cle_brute = match.group(1).strip().lower()
    valeur = match.group(2).strip()

    for champ, alias_list in LABELS.items():
        if cle_brute in alias_list:
            return champ, valeur

    return None


# DÉCOUPAGE DU TEXTE EN BLOCS (1 bloc = 1 sujet potentiel)

def _decouper_en_blocs(lignes: list) -> list:
    """Regroupe les lignes en blocs, séparés par une ligne vide ou par
    l'apparition d'un nouveau label 'Titre :' (utile pour les PDF compacts
    sans lignes vides entre les sujets)."""

    blocs = []
    bloc_courant = []

    for ligne in lignes:
        ligne_nettoyee = ligne.strip()

        if not ligne_nettoyee:
            if bloc_courant:
                blocs.append(bloc_courant)
                bloc_courant = []
            continue

        detection = _detecter_label(ligne_nettoyee)

        if detection and detection[0] == "titre" and bloc_courant:
            blocs.append(bloc_courant)
            bloc_courant = []

        bloc_courant.append(ligne_nettoyee)

    if bloc_courant:
        blocs.append(bloc_courant)

    return blocs


def _extraire_sujet_depuis_bloc(lignes_bloc: list) -> dict:
    sujet = {champ: None for champ in CHAMPS_TEXTE}
    sujet["titre"] = ""
    sujet["description"] = ""
    sujet["domaine_indique"] = None

    lignes_libres = []

    for ligne in lignes_bloc:
        detection = _detecter_label(ligne)

        if detection:
            champ, valeur = detection

            if champ == "domaine":
                sujet["domaine_indique"] = valeur
            elif sujet.get(champ):
                sujet[champ] += " " + valeur
            else:
                sujet[champ] = valeur

            continue

        lignes_libres.append(ligne)

    if not sujet["titre"]:
        # Pas de label "Titre :" trouvé : on prend la première ligne du bloc
        # en retirant une éventuelle numérotation ("1.", "1)", "- ", "•").
        premiere = lignes_libres.pop(0) if lignes_libres else "Sujet sans titre"
        sujet["titre"] = re.sub(r"^\s*(\d+[\.\)]|[-*•])\s*", "", premiere).strip()

    if lignes_libres:
        reste = " ".join(lignes_libres)
        sujet["description"] = (
            (sujet["description"] + " " + reste).strip()
            if sujet["description"] else reste
        )

    for champ in CHAMPS_TEXTE:
        if sujet.get(champ):
            sujet[champ] = _normaliser(sujet[champ])

    return sujet


def extraire_sujets_depuis_texte(texte: str) -> list:
    lignes = texte.splitlines()
    blocs = _decouper_en_blocs(lignes)

    sujets = [_extraire_sujet_depuis_bloc(bloc) for bloc in blocs]

    # On ignore les blocs trop courts pour être un vrai sujet
    # (en-têtes de page, mentions légales isolées, etc.).
    return [s for s in sujets if s["titre"] and len(s["titre"]) > 3]


# LECTURE DES FICHIERS PAR FORMAT

def _ocr_pdf(contenu: bytes) -> str:
    """Convertit chaque page en image et lance la reconnaissance de texte
    (Tesseract). Utilisé en secours pour les PDF scannés, sans texte
    sélectionnable."""

    from pdf2image import convert_from_bytes
    import pytesseract

    pages = convert_from_bytes(contenu)
    return "\n".join(pytesseract.image_to_string(page, lang="fra+eng") for page in pages)


def lire_pdf(contenu: bytes) -> list:
    lecteur = PdfReader(BytesIO(contenu))
    texte_total = "\n".join(page.extract_text() or "" for page in lecteur.pages)

    # Si l'extraction directe ne renvoie presque rien, le PDF est
    # probablement un scan : on retente en OCR.
    if len(texte_total.strip()) < 20 * len(lecteur.pages):
        try:
            texte_ocr = _ocr_pdf(contenu)
            if len(texte_ocr.strip()) > len(texte_total.strip()):
                texte_total = texte_ocr
        except Exception:
            pass

    return extraire_sujets_depuis_texte(texte_total)


def lire_docx(contenu: bytes) -> list:
    document = Document(BytesIO(contenu))
    lignes = [paragraphe.text for paragraphe in document.paragraphs]
    return extraire_sujets_depuis_texte("\n".join(lignes))


def _lignes_tabulaires_vers_sujets(lignes: list) -> list:
    """Convertit une liste de lignes (en-tête + données) en sujets, en
    reconnaissant les colonnes par leur intitulé. Utilisé par Excel et CSV."""

    if not lignes:
        return []

    entetes = [str(cellule).strip().lower() if cellule not in (None, "") else "" for cellule in lignes[0]]

    correspondance = {}
    for index, entete in enumerate(entetes):
        for champ, alias_list in LABELS.items():
            if entete in alias_list:
                correspondance[champ] = index
                break

    sujets = []

    for ligne in lignes[1:]:
        if not any(ligne):
            continue

        def valeur(champ, ligne=ligne):
            index = correspondance.get(champ)
            if index is None or index >= len(ligne):
                return None
            cellule = ligne[index]
            return str(cellule).strip() if cellule not in (None, "") else None

        titre = valeur("titre")
        if not titre:
            continue

        sujets.append({
            "titre": titre,
            "description": valeur("description") or "",
            "entreprise": valeur("entreprise"),
            "technologies": valeur("technologies"),
            "duree": valeur("duree"),
            "localisation": valeur("localisation"),
            "niveau_requis": valeur("niveau_requis"),
            "domaine_indique": valeur("domaine"),
        })

    return sujets


def lire_excel(contenu: bytes) -> list:
    """Lit un fichier Excel où chaque ligne représente un sujet, avec des
    colonnes reconnues par leur en-tête (insensible à la casse)."""

    classeur = load_workbook(BytesIO(contenu), data_only=True)
    feuille = classeur.active

    lignes = list(feuille.iter_rows(values_only=True))
    return _lignes_tabulaires_vers_sujets(lignes)


def lire_csv(contenu: bytes) -> list:
    """Lit un CSV où chaque ligne représente un sujet (même logique que
    l'Excel). Format le plus courant pour un export de base externe."""

    texte = contenu.decode("utf-8-sig", errors="replace")
    lecteur = csv.reader(StringIO(texte))
    lignes = list(lecteur)
    return _lignes_tabulaires_vers_sujets(lignes)


def lire_json(contenu: bytes) -> list:
    """Lit un export JSON (liste d'objets) — typiquement une API ou une
    base de données externe. Les clés doivent correspondre aux noms des
    champs (titre, description, entreprise, technologies, duree,
    localisation, niveau_requis)."""

    donnees = json.loads(contenu.decode("utf-8"))

    if isinstance(donnees, dict):
        donnees = donnees.get("sujets") or donnees.get("data") or [donnees]

    sujets = []

    for entree in donnees:
        if not isinstance(entree, dict) or not entree.get("titre"):
            continue

        sujets.append({
            "titre": str(entree.get("titre")).strip(),
            "description": str(entree.get("description") or "").strip(),
            "entreprise": entree.get("entreprise"),
            "technologies": entree.get("technologies"),
            "duree": entree.get("duree"),
            "localisation": entree.get("localisation"),
            "niveau_requis": entree.get("niveau_requis"),
            "domaine_indique": entree.get("domaine") or entree.get("departement"),
        })

    return sujets


def lire_fichier(nom_fichier: str, contenu: bytes) -> list:
    extension = nom_fichier.lower().rsplit(".", 1)[-1] if "." in nom_fichier else ""

    if extension == "pdf":
        return lire_pdf(contenu)
    if extension == "docx":
        return lire_docx(contenu)
    if extension in ("xlsx", "xls"):
        return lire_excel(contenu)
    if extension == "csv":
        return lire_csv(contenu)
    if extension == "json":
        return lire_json(contenu)

    raise ValueError(
        "Format de fichier non pris en charge. "
        "Formats acceptés : PDF (.pdf), Word (.docx), Excel (.xlsx), "
        "CSV (.csv) et JSON (.json)."
    )


# CATÉGORISATION AUTOMATIQUE PAR DOMAINE

CATEGORIES = {
    "Développement Web": [
        "web", "react", "angular", "vue", "html", "css", "javascript",
        "frontend", "front-end", "backend", "back-end", "fullstack",
        "django", "fastapi", "node", "php", "laravel",
    ],
    "Data / Intelligence Artificielle": [
        "data", "données", "donnees", "ia", "intelligence artificielle",
        "machine learning", "deep learning", "nlp", "spacy", "transformers",
        "tensorflow", "pytorch", "data science", "big data", "analytics",
    ],
    "Mobile": [
        "mobile", "android", "ios", "flutter", "react native", "swift", "kotlin",
    ],
    "Réseaux & Systèmes": [
        "réseau", "reseau", "systèmes", "systemes", "cisco", "linux",
        "administration système", "infrastructure", "cloud", "devops",
        "docker", "kubernetes",
    ],
    "Cybersécurité": [
        "sécurité", "securite", "cybersécurité", "cybersecurite", "pentest",
        "firewall", "cryptographie",
    ],
}


def categoriser_sujet(sujet: dict) -> str:
    texte = " ".join([
        sujet.get("titre") or "",
        sujet.get("description") or "",
        sujet.get("technologies") or "",
        sujet.get("domaine_indique") or "",
    ]).lower()

    meilleure_categorie = "Autre"
    meilleur_score = 0

    for categorie, mots_cles in CATEGORIES.items():
        score = sum(1 for mot in mots_cles if mot in texte)
        if score > meilleur_score:
            meilleur_score = score
            meilleure_categorie = categorie

    return meilleure_categorie


# DÉTECTION DE DOUBLONS

def detecter_doublon(titre_nouveau: str, sujets_existants: list, seuil: float = 0.80):
    """Compare le titre d'un sujet importé à ceux déjà en base et renvoie
    le sujet existant le plus proche si la similarité dépasse le seuil."""

    titre_normalise = _normaliser(titre_nouveau).lower()
    meilleur = None
    meilleur_score = 0.0

    for sujet_existant in sujets_existants:
        titre_existant = _normaliser(sujet_existant.titre or "").lower()
        score = SequenceMatcher(None, titre_normalise, titre_existant).ratio()

        if score > meilleur_score:
            meilleur_score = score
            meilleur = sujet_existant

    if meilleur and meilleur_score >= seuil:
        return meilleur, round(meilleur_score * 100)

    return None, 0


# DÉTECTION DE SUJETS OBSOLÈTES

def est_obsolete(sujet, seuil_jours: int = SEUIL_OBSOLESCENCE_JOURS) -> bool:
    """Un sujet est considéré obsolète s'il est toujours "Disponible"
    (jamais attribué) alors qu'il a été créé il y a plus de seuil_jours."""

    statut = (sujet.statut or "").strip().lower()
    if statut != "disponible":
        return False

    date_creation = getattr(sujet, "date_creation", None)
    if not date_creation:
        return False

    return (date.today() - date_creation) > timedelta(days=seuil_jours)
